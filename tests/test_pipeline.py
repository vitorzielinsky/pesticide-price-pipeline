import os
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from src.catalog import CatalogValidationError, load_catalog
import src.database as db_module
import src.exporter as exporter_module
from src.extractor import extract_prices
from src.parser import parse_report_html
from src.scraper import (
    HttpResponse,
    SiteTemporarilyUnavailable,
    collect_listing_ids,
)
from src.review_importer import ReviewImportError, read_manual_decisions


class ExtractorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.catalog = load_catalog()

    def extract(self, text):
        return extract_prices(
            text,
            "https://example.test/report/1",
            "2026-06-09",
            1,
            self.catalog,
        )

    def test_units_with_spaces_and_full_width_slash(self):
        result = self.extract("草铵膦原药报价 5 万 元 ／ 吨；")
        self.assertEqual(len(result.prices), 1)
        self.assertEqual(result.prices[0]["preco_yuan_ton"], 50000)

    def test_wan_conversion_has_no_floating_point_residue(self):
        result = self.extract("二甲四氯原药报2.28万元/吨；")
        self.assertEqual(result.prices[0]["preco_yuan_ton"], 22800.0)

    def test_explicit_concentration_has_priority(self):
        result = self.extract("草甘膦原药95%含量报3.55万元/吨；")
        self.assertEqual(result.prices[0]["concentracao"], "95%")
        self.assertEqual(result.prices[0]["preco_yuan_ton"], 35500)

    def test_base_100_has_priority(self):
        result = self.extract("精草铵膦原药折百报价6.5万元/吨；")
        self.assertEqual(result.prices[0]["produto_chines"], "精草铵膦")
        self.assertEqual(result.prices[0]["concentracao"], "100% (折百)")

    def test_unknown_product_goes_to_review(self):
        result = self.extract("未知中间体报价5.5万元/吨；")
        self.assertEqual(result.prices, [])
        self.assertEqual(result.reviews[0]["produto_sugerido"], "produto_desconhecido")

    def test_bentazone_technical_and_liquid_forms_are_separated(self):
        result = self.extract(
            "灭草松原药出口为主，"
            "报价7.5万元/吨；"
            "480克/升灭草松水剂"
            "报价3万元/吨；"
        )
        by_price = {price["preco_yuan_ton"]: price for price in result.prices}
        self.assertEqual(set(by_price), {75000, 30000})
        self.assertEqual(by_price[75000]["produto_chines"], "灭草松")
        self.assertEqual(by_price[75000]["forma_preco"], "tecnico_original")
        self.assertIn("灭草松原药", by_price[75000]["forma_original"])
        self.assertEqual(by_price[30000]["produto_chines"], "灭草松")
        self.assertEqual(by_price[30000]["forma_preco"], "liquido_agua")
        self.assertEqual(by_price[30000]["concentracao"], "480克/升")
        self.assertIn("灭草松水剂", by_price[30000]["forma_original"])

    def test_bentazone_liquid_mention_does_not_contaminate_technical_quote(self):
        result = self.extract(
            "灭草松原药供应紧张，"
            "厂家生产水剂为主，"
            "原药报到14.5万元/吨，"
            "灭草松水剂报到4.8万元/吨；"
        )
        by_price = {price["preco_yuan_ton"]: price for price in result.prices}
        self.assertEqual(by_price[145000]["forma_preco"], "tecnico_original")
        self.assertEqual(by_price[48000]["forma_preco"], "liquido_agua")

    def test_quote_at_clause_start_uses_previous_product_context(self):
        result = self.extract(
            "灭草松原药厂家环保压力"
            "较大，原药出口为主，"
            "价格有所上调。"
            "报到12万元/吨左右，"
            "灭草松水剂市场上调"
            "价格至4.3万元/吨；"
        )
        by_price = {price["preco_yuan_ton"]: price for price in result.prices}
        self.assertEqual(by_price[120000]["produto_chines"], "灭草松")
        self.assertEqual(by_price[120000]["forma_preco"], "tecnico_original")
        self.assertEqual(
            by_price[120000]["regra_aplicada"],
            "contexto_frase_anterior",
        )

    def test_pendimethalin_does_not_receive_mcpa_price(self):
        result = self.extract(
            "二甲戊灵原药报5.1万元/吨，"
            "二甲四氯钠盐供应紧张，"
            "报2.1万元/吨以上；"
        )
        by_price = {price["preco_yuan_ton"]: price for price in result.prices}
        self.assertEqual(by_price[51000]["produto_chines"], "二甲戊灵")
        self.assertEqual(by_price[21000]["produto_chines"], "二甲四氯")
        self.assertEqual(by_price[21000]["forma_preco"], "sal_sodico")

    def test_abamectin_and_emamectin_are_separated(self):
        result = self.extract(
            "阿维菌素精粉报价75万元/吨，"
            "甲维盐原药报价125万元/吨；"
        )
        by_price = {price["preco_yuan_ton"]: price for price in result.prices}
        self.assertEqual(by_price[750000]["produto_chines"], "阿维菌素")
        self.assertEqual(by_price[750000]["forma_preco"], "po_cristalino")
        self.assertEqual(
            by_price[1250000]["produto_chines"],
            "甲氨基阿维菌素苯甲酸盐",
        )

    def test_chinese_suffix_match_goes_to_review(self):
        result = self.extract(
            "唑醚菌酯原药市场报价26"
            "万元/吨；"
            "醚菌酯原药价格报35"
            "万元/吨；"
        )
        self.assertEqual(len(result.prices), 1)
        self.assertEqual(result.prices[0]["produto_chines"], "醚菌酯")
        self.assertEqual(result.prices[0]["preco_yuan_ton"], 350000)
        self.assertEqual(len(result.reviews), 1)
        self.assertEqual(result.reviews[0]["produto_sugerido"], "produto_desconhecido")
        self.assertEqual(result.reviews[0]["preco_yuan_ton"], 260000)

    def test_24d_and_mcpa_stay_separate(self):
        result = self.extract(
            "二甲四氯钠盐报价2.1"
            "万元/吨；"
            "2,4-D原药报价3万元/吨；"
        )
        by_price = {price["preco_yuan_ton"]: price for price in result.prices}
        self.assertEqual(by_price[21000]["produto_chines"], "二甲四氯")
        self.assertEqual(by_price[30000]["produto_chines"], "2,4-D")

    def test_24d_with_spaces_around_punctuation_is_recognized(self):
        result = self.extract(
            "2 ,4-D原药行情平稳，"
            "本月价格报到2.4万元/吨；"
        )
        self.assertEqual(result.reviews, [])
        self.assertEqual(result.prices[0]["produto_chines"], "2,4-D")
        self.assertEqual(result.prices[0]["preco_yuan_ton"], 24000)

    def test_unknown_product_near_quote_is_not_inherited(self):
        result = self.extract(
            "马拉硫磷原药报价1.95"
            "万元/吨，"
            "三唑磷报价4.2万元/吨。"
        )
        self.assertEqual(len(result.prices), 1)
        self.assertEqual(result.prices[0]["produto_chines"], "马拉硫磷")
        self.assertEqual(result.prices[0]["preco_yuan_ton"], 19500)
        self.assertEqual(len(result.reviews), 1)
        self.assertEqual(result.reviews[0]["produto_sugerido"], "produto_desconhecido")
        self.assertEqual(result.reviews[0]["preco_yuan_ton"], 42000)

    def test_unknown_product_after_previous_quote_is_not_inherited(self):
        result = self.extract(
            "马拉硫磷原药开工率低，"
            "厂家不愿接单，报价1.95"
            "万元/吨，三唑磷开工率"
            "降低，货源紧张，报价4.2"
            "万元/吨。"
        )
        self.assertEqual(len(result.prices), 1)
        self.assertEqual(result.prices[0]["produto_chines"], "马拉硫磷")
        self.assertEqual(result.prices[0]["preco_yuan_ton"], 19500)
        self.assertEqual(len(result.reviews), 1)
        self.assertEqual(result.reviews[0]["produto_sugerido"], "produto_desconhecido")
        self.assertEqual(result.reviews[0]["preco_yuan_ton"], 42000)

    def test_spaced_chinese_aliases_are_recognized(self):
        cases = [
            (
                "茚虫 威原药价格报到95"
                "万元/吨；",
                "茚虫威",
                950000,
            ),
            (
                "氟 吡菌 胺原 药市场"
                "价格报到85万元/吨；",
                "氟吡菌胺",
                850000,
            ),
            (
                "阿维菌素精 粉报到45"
                "万元/吨；",
                "阿维菌素",
                450000,
            ),
        ]
        for text, product, price in cases:
            with self.subTest(product=product):
                result = self.extract(text)
                self.assertEqual(result.reviews, [])
                self.assertEqual(len(result.prices), 1)
                self.assertEqual(result.prices[0]["produto_chines"], product)
                self.assertEqual(result.prices[0]["preco_yuan_ton"], price)

    def test_form_original_is_normalized_without_internal_spaces(self):
        result = self.extract(
            "氟磺胺 草醚原 药"
            "市场报到14万元/吨；"
        )
        self.assertEqual(result.reviews, [])
        self.assertEqual(result.prices[0]["produto_chines"], "氟磺胺草醚")
        self.assertEqual(
            result.prices[0]["forma_original"],
            "氟磺胺草醚原药",
        )

    def test_operational_words_after_known_product_do_not_force_review(self):
        result = self.extract(
            "甲氨基阿维菌素苯甲酸"
            "盐原药工厂报价64"
            "万元/吨；"
        )
        self.assertEqual(result.reviews, [])
        self.assertEqual(len(result.prices), 1)
        self.assertEqual(
            result.prices[0]["produto_chines"],
            "甲氨基阿维菌素苯甲酸盐",
        )
        self.assertEqual(result.prices[0]["preco_yuan_ton"], 640000)

    def test_operational_context_keeps_known_product(self):
        result = self.extract(
            "烯草酮原药原材料"
            "快速上涨，当前报价8"
            "万元/吨；"
        )
        self.assertEqual(result.reviews, [])
        self.assertEqual(len(result.prices), 1)
        self.assertEqual(result.prices[0]["produto_chines"], "烯草酮")
        self.assertEqual(result.prices[0]["preco_yuan_ton"], 80000)

    def test_temporal_context_before_quote_keeps_known_product(self):
        cases = [
            (
                "烯草酮原药上游中间体"
                "上涨，本周报价12万元/吨；",
                "烯草酮",
                120000,
            ),
            (
                "嘧菌酯原药供应紧张，"
                "本周报到35万元/吨；",
                "嘧菌酯",
                350000,
            ),
        ]
        for text, product, price in cases:
            with self.subTest(product=product):
                result = self.extract(text)
                self.assertEqual(result.reviews, [])
                self.assertEqual(len(result.prices), 1)
                self.assertEqual(result.prices[0]["produto_chines"], product)
                self.assertEqual(result.prices[0]["preco_yuan_ton"], price)

    def test_context_prefix_before_chinese_alias_is_allowed(self):
        cases = [
            (
                "目前高效氯氟氰菊酯原药"
                "市场报到12.8万元/吨；",
                "高效氯氟氰菊酯",
                128000,
            ),
            (
                "有机磷类产品马拉硫磷原药"
                "市场报到3.5万元/吨；",
                "马拉硫磷",
                35000,
            ),
        ]
        for text, product, price in cases:
            with self.subTest(product=product):
                result = self.extract(text)
                self.assertEqual(result.reviews, [])
                self.assertEqual(result.prices[0]["produto_chines"], product)
                self.assertEqual(result.prices[0]["preco_yuan_ton"], price)

    def test_context_after_chinese_alias_is_allowed(self):
        cases = [
            (
                "精甲霜灵市场报价13.6"
                "万元/吨；",
                "精甲霜灵",
                136000,
            ),
            (
                "草甘膦受原材料价格"
                "影响，报到6.55万元/吨；",
                "草甘膦",
                65500,
            ),
        ]
        for text, product, price in cases:
            with self.subTest(product=product):
                result = self.extract(text)
                self.assertEqual(result.reviews, [])
                self.assertEqual(result.prices[0]["produto_chines"], product)
                self.assertEqual(result.prices[0]["preco_yuan_ton"], price)

    def test_parallel_product_price_list_is_paired_by_position(self):
        result = self.extract(
            "多菌灵原药、甲基硫菌灵原药"
            "供货稳定，报价3.8万元/吨，"
            "报价3万元/吨。"
        )
        by_price = {price["preco_yuan_ton"]: price for price in result.prices}
        self.assertEqual(by_price[38000]["produto_chines"], "多菌灵")
        self.assertEqual(by_price[30000]["produto_chines"], "甲基硫菌灵")

    def test_respective_product_price_list_is_paired_by_position(self):
        result = self.extract(
            "吡虫啉原药、啶虫脒原药"
            "分别报18.8万元/吨，19.2万元/吨；"
        )
        by_price = {price["preco_yuan_ton"]: price for price in result.prices}
        self.assertEqual(by_price[188000]["produto_chines"], "吡虫啉")
        self.assertEqual(by_price[192000]["produto_chines"], "啶虫脒")

    def test_shared_price_is_applied_to_connected_products(self):
        cases = [
            "吡虫啉原药和啶虫脒原药"
            "报价17万元/吨；",
            "吡虫啉和啶虫脒原药"
            "报17万元/吨；",
        ]
        for text in cases:
            with self.subTest(text=text):
                result = self.extract(text)
                found = {
                    (price["produto_chines"], price["preco_yuan_ton"], price["regra_aplicada"])
                    for price in result.prices
                }
                self.assertEqual(result.reviews, [])
                self.assertIn(
                    ("吡虫啉", 170000, "preco_compartilhado"),
                    found,
                )
                self.assertIn(
                    ("啶虫脒", 170000, "preco_compartilhado"),
                    found,
                )

    def test_intermediate_name_does_not_hide_chlorpyrifos_quote(self):
        result = self.extract(
            "丙溴磷原药价格保持在6.4万元/吨，"
            "毒死蜱原药成本不断抬升，"
            "主要利润流向三氯吡啶醇钠生产厂家，"
            "本月报到3.62万元/吨，"
            "辛硫磷原药当前报2.71万元/吨。"
        )
        by_price = {price["preco_yuan_ton"]: price for price in result.prices}
        self.assertEqual(by_price[36200]["produto_chines"], "毒死蜱")
        self.assertEqual(by_price[27100]["produto_chines"], "辛硫磷")

    def test_actual_thiamethoxam_seed_quality_context_is_not_unknown(self):
        result = self.extract(
            "噻虫嗪原药种衣剂品质"
            "报到12万元/吨，市场流通"
            "货源报10.5万元/吨；"
        )
        by_price = {price["preco_yuan_ton"]: price for price in result.prices}
        self.assertEqual(by_price[120000]["segmento_preco"], "qualidade_tratamento_sementes")
        self.assertEqual(by_price[105000]["segmento_preco"], "mercado_circulante")

    def test_explicit_and_qualitative_concentrations_are_captured(self):
        explicit = self.extract(
            "噻虫嗪原药部分拌种剂需用"
            "98%含量货源，目前价格报到"
            "9.2万元/吨；"
        )
        self.assertEqual(explicit.prices[0]["concentracao"], "98%")

        qualitative = self.extract(
            "异噁草松原药国内市场"
            "低含量货源价格报到8万元/吨；"
        )
        self.assertEqual(qualitative.prices[0]["concentracao"], "baixo_teor")

    def test_diquat_salt_forms_are_separated(self):
        result = self.extract(
            "敌草快母药国内市场"
            "二氯盐报到4.2万元/吨，"
            "二溴盐报到4.4万元/吨。"
        )
        by_price = {price["preco_yuan_ton"]: price for price in result.prices}
        self.assertEqual(by_price[42000]["forma_preco"], "sal_dicloreto")
        self.assertEqual(by_price[44000]["forma_preco"], "sal_dibrometo")

    def test_quality_and_licensed_factory_contexts_are_not_products(self):
        cases = [
            (
                "噻唑膦原药高端原药"
                "报价29万元/吨；",
                290000,
                "alta_qualidade",
            ),
            (
                "环磺酮原药有证大厂"
                "报价32万元/吨；",
                320000,
                "canal_fabrica",
            ),
        ]
        for text, price, segment in cases:
            with self.subTest(segment=segment):
                result = self.extract(text)
                self.assertEqual(result.reviews, [])
                self.assertEqual(result.prices[0]["preco_yuan_ton"], price)
                self.assertEqual(result.prices[0]["segmento_preco"], segment)

    def test_fomesafen_broad_concentration_pattern(self):
        result = self.extract(
            "氟磺胺草醚原药95%含量"
            "货源市场报到14万元/吨，"
            "97%含量货源市场报到14.5"
            "万元/吨；"
        )
        found = {
            (price["produto_chines"], price["concentracao"], price["preco_yuan_ton"])
            for price in result.prices
        }
        self.assertIn(("氟磺胺草醚", "95%", 140000), found)
        self.assertIn(("氟磺胺草醚", "97%", 145000), found)

    def test_thiamethoxam_segments_are_separated(self):
        result = self.extract(
            "噻虫嗪原药种衣剂品质"
            "产品报到12万元/吨，"
            "市场流通货源报10.5"
            "万元/吨；"
        )
        by_price = {price["preco_yuan_ton"]: price for price in result.prices}
        self.assertEqual(
            by_price[120000]["segmento_preco"],
            "qualidade_tratamento_sementes",
        )
        self.assertEqual(
            by_price[105000]["segmento_preco"],
            "mercado_circulante",
        )

    def test_concentration_variants_are_captured(self):
        result = self.extract(
            "草甘膦原药95%含量暂报"
            "2.75万元/吨，97%含量报到"
            "2.8万元/吨；"
            "精草铵膦原药折百报"
            "6.5万元/吨；"
        )
        found = {
            (price["produto_chines"], price["concentracao"], price["preco_yuan_ton"])
            for price in result.prices
        }
        self.assertIn(("草甘膦", "95%", 27500), found)
        self.assertIn(("草甘膦", "97%", 28000), found)
        self.assertIn(("精草铵膦", "100% (折百)", 65000), found)

    def test_report_really_without_quote_returns_empty(self):
        result = self.extract("Relatório informativo sem valores de mercado.")
        self.assertEqual(result.prices, [])
        self.assertEqual(result.reviews, [])
        self.assertEqual(result.quote_count, 0)


class CatalogTests(unittest.TestCase):
    def _workbook(self, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Herbicidas"
        sheet.append(["Título", None, None])
        sheet.append(["Explicação", None, None])
        sheet.append(["Português", "Inglês", "Chinês"])
        for row in rows:
            sheet.append(row)
        for name in ("Inseticidas", "Fungicidas"):
            extra_sheet = workbook.create_sheet(name)
            extra_sheet.append(["Português", "Inglês", "Chinês"])
        other_sheet = workbook.create_sheet("Outros do Banco")
        other_sheet.append([
            "Português", "Inglês", "Chinês", "Categoria",
        ])
        return workbook

    def test_missing_catalog_stops_execution(self):
        with self.assertRaises(CatalogValidationError):
            load_catalog("arquivo_que_nao_existe.xlsx")

    def test_duplicate_chinese_name_stops_execution(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            path = os.path.join(directory, "catalog.xlsx")
            workbook = self._workbook([
                ["Produto A", "Product A", "中文名"],
                ["Produto A", "Product A", "中文名"],
            ])
            workbook.save(path)
            workbook.close()
            with self.assertRaises(CatalogValidationError):
                load_catalog(path)

    def test_empty_required_field_stops_execution(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            path = os.path.join(directory, "catalog.xlsx")
            workbook = self._workbook([
                ["Produto A", "", "中文名"],
            ])
            workbook.save(path)
            workbook.close()
            with self.assertRaises(CatalogValidationError):
                load_catalog(path)

    def test_missing_required_sheet_stops_execution(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            path = os.path.join(directory, "catalog.xlsx")
            workbook = Workbook()
            workbook.active.title = "Herbicidas"
            workbook.active.append(["Português", "Inglês", "Chinês"])
            workbook.save(path)
            workbook.close()
            with self.assertRaises(CatalogValidationError) as context:
                load_catalog(path)
            self.assertIn("aba", str(context.exception))


class ParserTests(unittest.TestCase):
    def test_real_title_is_taken_from_article(self):
        html = """
        <html><head><title>Site</title></head><body>
          <header><h3>产品中心</h3></header>
          <div class="news-content">
            <h1>Título verdadeiro</h1>
            <p>发布时间：2026-06-09</p><p>报价5万元/吨</p>
          </div>
        </body></html>
        """
        result = parse_report_html(
            html,
            "https://example.test/report?id=10",
            "https://example.test/report?id=10",
            10,
        )
        self.assertEqual(result["titulo"], "Título verdadeiro")

    def test_service_menu_is_not_used_as_title(self):
        html = """
        <html><body>
          <h3>产品中心</h3><h3>首页</h3><h3>新闻中心</h3>
          <div class="content">
            <h3>Título verdadeiro do relatório</h3>
            <p>发布时间：2026-06-09</p><p>报价5万元/吨</p>
          </div>
        </body></html>
        """
        result = parse_report_html(
            html,
            "https://example.test/report?id=10",
            "https://example.test/report?id=10",
            10,
        )
        self.assertEqual(
            result["titulo"], "Título verdadeiro do relatório"
        )

    def test_redirected_id_is_rejected(self):
        html = "<html><body><div class='content'>texto</div></body></html>"
        with self.assertRaises(ValueError):
            parse_report_html(
                html,
                "https://example.test/report?id=10",
                "https://example.test/report?id=11",
                10,
            )


class ScraperTests(unittest.TestCase):
    def setUp(self):
        for name, value in (
            ("src.scraper.LIST_URL", "https://example.test/list?page={page}"),
            ("src.scraper.REPORT_URL", "https://example.test/report?id={id}"),
        ):
            patcher = patch(name, value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def test_first_page_failure_is_reported_as_site_unavailable(self):
        with patch(
            "src.scraper.safe_request",
            return_value=None,
        ):
            with self.assertRaises(SiteTemporarilyUnavailable):
                collect_listing_ids()

    def test_dynamic_pagination_stops_on_empty_page(self):
        page_one = b"""
        <html><body>
          <div><a href="report?id=10">Relatorio</a>
          <em>2026-06-09</em></div>
        </body></html>
        """
        empty_page = b"<html><body>Sem relatorios</body></html>"
        responses = [
            HttpResponse(page_one.decode(), "https://example/1", 200, page_one),
            HttpResponse(empty_page.decode(), "https://example/2", 200, empty_page),
        ]
        with patch(
            "src.scraper.safe_request",
            side_effect=responses,
        ), patch(
            "src.scraper.report_exists", return_value=False
        ), patch("src.scraper.time.sleep"):
            result = collect_listing_ids(stop_after_known_page=False)
        self.assertEqual(
            result.items, [(10, "2026-06-09", "Relatorio")]
        )
        self.assertTrue(result.complete)

    def test_update_stops_after_page_with_only_known_reports(self):
        page_one = b"""
        <html><body>
          <div><a href="report?id=10">Relatorio</a>
          <em>2026-06-09</em></div>
        </body></html>
        """
        response = HttpResponse(
            page_one.decode(), "https://example/1", 200, page_one
        )
        with patch(
            "src.scraper.safe_request",
            return_value=response,
        ) as request_mock, patch(
            "src.scraper.report_exists",
            return_value=True,
        ), patch("src.scraper.time.sleep"):
            result = collect_listing_ids(stop_after_known_page=True)

        self.assertEqual(
            result.items, [(10, "2026-06-09", "Relatorio")]
        )
        self.assertTrue(result.complete)
        self.assertEqual(request_mock.call_count, 1)

    def test_later_page_failure_returns_partial_results(self):
        page_one = b"""
        <html><body>
          <div><a href="report?id=11">Novo</a>
          <em>2026-06-09</em></div>
        </body></html>
        """
        responses = [
            HttpResponse(page_one.decode(), "https://example/1", 200, page_one),
            None,
        ]
        with patch(
            "src.scraper.safe_request",
            side_effect=responses,
        ), patch(
            "src.scraper.report_exists",
            return_value=False,
        ), patch("src.scraper.time.sleep"):
            result = collect_listing_ids(stop_after_known_page=True)

        self.assertEqual(result.items, [(11, "2026-06-09", "Novo")])
        self.assertFalse(result.complete)
        self.assertEqual(result.failed_page, 2)


class DatabaseTests(unittest.TestCase):
    def test_intermediate_market_quote_is_not_suggested_as_downstream_product(self):
        catalog = load_catalog()
        status, match = db_module._review_catalog_assessment(
            catalog,
            (
                "中间体市场 CCMP 厂家意欲"
                "上调价格，下游不少"
                "吡虫啉装置计划开车生产，"
                "报到10万元/吨以上；"
            ),
            100000,
        )
        self.assertEqual(status, "none")
        self.assertIsNone(match)

    def test_reprocessing_replaces_instead_of_duplicating(self):
        with tempfile.TemporaryDirectory() as directory:
            database_path = os.path.join(directory, "corrected.db")
            with patch.object(
                db_module, "DB_PATH", database_path
            ):
                db_module.initialize_database()
                db_module.upsert_report({
                    "id_relatorio": 1,
                    "url_solicitada": "https://example.test/1",
                    "url_final": "https://example.test/1",
                    "codigo_http": 200,
                    "data_relatorio": "2026-06-09",
                    "titulo": "Teste",
                    "texto_completo": "草铵膦报价5万元/吨",
                    "html_bruto": "<p>teste</p>",
                    "hash_conteudo": "abc",
                    "eh_mercado": True,
                })
                price = {
                    "data_relatorio": "2026-06-09",
                    "produto_chines": "草铵膦",
                    "produto_ingles": "Glufosinate-ammonium",
                    "produto_portugues": "Glufosinato de amônio",
                    "categoria": "herbicida",
                    "concentracao": "",
                    "preco_yuan_ton": 50000,
                    "unidade_original": "5万元/吨",
                    "url": "https://example.test/1",
                    "trecho_original": "草铵膦报价5万元/吨",
                    "nome_encontrado": "草铵膦",
                    "posicao_texto": 4,
                    "regra_aplicada": "cotacao_generica",
                    "confianca": "alta",
                    "extractor_version": "2.0.0",
                    "catalog_version": "teste",
                }
                db_module.replace_report_results(
                    1, [price], [], "extraido"
                )
                db_module.replace_report_results(
                    1, [price], [], "extraido"
                )
                connection = sqlite3.connect(database_path)
                count = connection.execute(
                    "SELECT COUNT(*) FROM precos"
                ).fetchone()[0]
                connection.close()
                self.assertEqual(count, 1)

    def test_manual_price_survives_automatic_reprocessing(self):
        with tempfile.TemporaryDirectory() as directory:
            database_path = os.path.join(directory, "corrected.db")
            with patch.object(
                db_module, "DB_PATH", database_path
            ):
                db_module.initialize_database()
                db_module.upsert_report({
                    "id_relatorio": 1,
                    "url_solicitada": "https://example.test/1",
                    "url_final": "https://example.test/1",
                    "codigo_http": 200,
                    "data_relatorio": "2026-06-09",
                    "titulo": "Teste",
                    "texto_completo": "texto",
                    "html_bruto": "<p>teste</p>",
                    "hash_conteudo": "abc",
                    "eh_mercado": True,
                })
                with db_module.connection() as connection:
                    connection.execute("""
                        INSERT INTO precos (
                            id_relatorio, data_relatorio, produto_chines,
                            produto_ingles, produto_portugues, categoria,
                            concentracao, preco_yuan_ton, unidade_original,
                            url, trecho_original, nome_encontrado,
                            posicao_texto, regra_aplicada, confianca,
                            extractor_version, catalog_version
                        ) VALUES (
                            1, '2026-06-09', '草铵膦',
                            'Glufosinate-ammonium', 'GLUFOSINATO',
                            'herbicida', '', 50000, 'Manual',
                            'https://example.test/1', 'trecho', '草铵膦',
                            -10, 'revisao_manual_aprovada', 'alta',
                            'manual', 'teste'
                        )
                    """)
                db_module.replace_report_results(
                    1, [], [], "sem_preco_confirmado"
                )
                connection = sqlite3.connect(database_path)
                count = connection.execute("""
                    SELECT COUNT(*) FROM precos
                    WHERE regra_aplicada = 'revisao_manual_aprovada'
                """).fetchone()[0]
                connection.close()
                self.assertEqual(count, 1)

    def test_automatic_review_is_removed_on_reprocessing(self):
        with tempfile.TemporaryDirectory() as directory:
            database_path = os.path.join(directory, "corrected.db")
            with patch.object(
                db_module, "DB_PATH", database_path
            ):
                db_module.initialize_database()
                db_module.upsert_report({
                    "id_relatorio": 1,
                    "url_solicitada": "https://example.test/1",
                    "url_final": "https://example.test/1",
                    "codigo_http": 200,
                    "data_relatorio": "2026-06-09",
                    "titulo": "Teste",
                    "texto_completo": "texto",
                    "html_bruto": "<p>teste</p>",
                    "hash_conteudo": "abc",
                    "eh_mercado": True,
                })
                db_module.replace_report_results(1, [], [{
                    "data_relatorio": "2026-06-09",
                    "produto_sugerido": "produto_desconhecido",
                    "preco_yuan_ton": 55000,
                    "trecho_original": "未知产品报5.5万元/吨",
                    "motivo": "teste",
                    "confianca": "baixa",
                }], "revisao_manual")
                db_module.resolve_unknown_reviews()
                db_module.replace_report_results(
                    1, [], [], "sem_preco_confirmado"
                )
                self.assertEqual(db_module.fetch_reviews(), [])

    def test_unknown_review_with_catalog_alias_stays_pending(self):
        catalog = load_catalog()
        with tempfile.TemporaryDirectory() as directory:
            database_path = os.path.join(directory, "corrected.db")
            with patch.object(
                db_module, "DB_PATH", database_path
            ):
                db_module.initialize_database()
                db_module.upsert_report({
                    "id_relatorio": 1,
                    "url_solicitada": "https://example.test/1",
                    "url_final": "https://example.test/1",
                    "codigo_http": 200,
                    "data_relatorio": "2026-06-09",
                    "titulo": "Teste",
                    "texto_completo": "texto",
                    "html_bruto": "<p>teste</p>",
                    "hash_conteudo": "abc",
                    "eh_mercado": True,
                })
                db_module.replace_report_results(1, [], [{
                    "data_relatorio": "2026-06-09",
                    "produto_sugerido": "produto_desconhecido",
                    "preco_yuan_ton": 640000,
                    "trecho_original": (
                        "甲氨基阿维菌素"
                        "苯甲酸盐原药工厂"
                        "报价64万元/吨"
                    ),
                    "motivo": "teste",
                    "confianca": "baixa",
                }], "revisao_manual")
                resolved = db_module.resolve_unknown_reviews(catalog)
                reviews = db_module.fetch_reviews()
                self.assertEqual(resolved, 0)
                self.assertEqual(reviews[0][3], "Emamectina")
                self.assertEqual(reviews[0][8], "")

    def test_auto_closed_unknown_with_catalog_alias_is_reopened(self):
        catalog = load_catalog()
        with tempfile.TemporaryDirectory() as directory:
            database_path = os.path.join(directory, "corrected.db")
            with patch.object(
                db_module, "DB_PATH", database_path
            ):
                db_module.initialize_database()
                db_module.upsert_report({
                    "id_relatorio": 1,
                    "url_solicitada": "https://example.test/1",
                    "url_final": "https://example.test/1",
                    "codigo_http": 200,
                    "data_relatorio": "2026-06-09",
                    "titulo": "Teste",
                    "texto_completo": "texto",
                    "html_bruto": "<p>teste</p>",
                    "hash_conteudo": "abc",
                    "eh_mercado": True,
                })
                db_module.replace_report_results(1, [], [{
                    "data_relatorio": "2026-06-09",
                    "produto_sugerido": "produto_desconhecido",
                    "preco_yuan_ton": 640000,
                    "trecho_original": (
                        "甲氨基阿维菌素"
                        "苯甲酸盐原药工厂"
                        "报价64万元/吨"
                    ),
                    "motivo": "teste",
                    "confianca": "baixa",
                }], "revisao_manual")
                db_module.resolve_unknown_reviews()
                db_module.resolve_unknown_reviews(catalog)
                reviews = db_module.fetch_reviews()
                self.assertEqual(reviews[0][3], "Emamectina")
                self.assertEqual(reviews[0][8], "")
                self.assertEqual(reviews[0][9], "")
                self.assertEqual(reviews[0][10], "")

    def test_unknown_review_uses_price_context_before_suggesting_alias(self):
        catalog = load_catalog()
        with tempfile.TemporaryDirectory() as directory:
            database_path = os.path.join(directory, "corrected.db")
            with patch.object(
                db_module, "DB_PATH", database_path
            ):
                db_module.initialize_database()
                db_module.upsert_report({
                    "id_relatorio": 1,
                    "url_solicitada": "https://example.test/1",
                    "url_final": "https://example.test/1",
                    "codigo_http": 200,
                    "data_relatorio": "2026-06-09",
                    "titulo": "Teste",
                    "texto_completo": "texto",
                    "html_bruto": "<p>teste</p>",
                    "hash_conteudo": "abc",
                    "eh_mercado": True,
                })
                db_module.replace_report_results(1, [], [{
                    "data_relatorio": "2026-06-09",
                    "produto_sugerido": "produto_desconhecido",
                    "preco_yuan_ton": 42000,
                    "trecho_original": (
                        "马拉硫磷原药报价1.95"
                        "万元/吨，三唑磷"
                        "货源紧张，报价4.2"
                        "万元/吨"
                    ),
                    "motivo": "teste",
                    "confianca": "baixa",
                }], "revisao_manual")
                resolved = db_module.resolve_unknown_reviews(catalog)
                reviews = db_module.fetch_reviews()
                self.assertEqual(resolved, 1)
                self.assertEqual(reviews[0][3], "produto_desconhecido")
                self.assertEqual(reviews[0][8], "N\u00e3o incluir")

    def test_unknown_review_without_catalog_alias_is_auto_closed(self):
        catalog = load_catalog()
        with tempfile.TemporaryDirectory() as directory:
            database_path = os.path.join(directory, "corrected.db")
            with patch.object(
                db_module, "DB_PATH", database_path
            ):
                db_module.initialize_database()
                db_module.upsert_report({
                    "id_relatorio": 1,
                    "url_solicitada": "https://example.test/1",
                    "url_final": "https://example.test/1",
                    "codigo_http": 200,
                    "data_relatorio": "2026-06-09",
                    "titulo": "Teste",
                    "texto_completo": "texto",
                    "html_bruto": "<p>teste</p>",
                    "hash_conteudo": "abc",
                    "eh_mercado": True,
                })
                db_module.replace_report_results(1, [], [{
                    "data_relatorio": "2026-06-09",
                    "produto_sugerido": "produto_desconhecido",
                    "preco_yuan_ton": 55000,
                    "trecho_original": (
                        "未知中间体报价5.5"
                        "万元/吨"
                    ),
                    "motivo": "teste",
                    "confianca": "baixa",
                }], "revisao_manual")
                resolved = db_module.resolve_unknown_reviews(catalog)
                reviews = db_module.fetch_reviews()
                self.assertEqual(resolved, 1)
                self.assertEqual(reviews[0][3], "produto_desconhecido")
                self.assertEqual(reviews[0][8], "Não incluir")

    def test_quality_checks_do_not_reopen_manual_unknown_decisions(self):
        catalog = load_catalog()
        with tempfile.TemporaryDirectory() as directory:
            database_path = os.path.join(directory, "corrected.db")
            with patch.object(
                db_module, "DB_PATH", database_path
            ):
                db_module.initialize_database()
                db_module.upsert_report({
                    "id_relatorio": 1,
                    "url_solicitada": "https://example.test/1",
                    "url_final": "https://example.test/1",
                    "codigo_http": 200,
                    "data_relatorio": "2026-06-09",
                    "titulo": "Teste",
                    "texto_completo": "texto",
                    "html_bruto": "<p>teste</p>",
                    "hash_conteudo": "abc",
                    "eh_mercado": True,
                })
                db_module.replace_report_results(1, [], [{
                    "data_relatorio": "2026-06-09",
                    "produto_sugerido": "produto_desconhecido",
                    "preco_yuan_ton": 260000,
                    "trecho_original": (
                        "唑醚菌酯原药"
                        "市场报价26万元/吨"
                    ),
                    "motivo": "teste",
                    "confianca": "baixa",
                }], "revisao_manual")
                with db_module.connection() as connection:
                    connection.execute("""
                        UPDATE revisao_manual
                        SET decisao = 'Não incluir',
                            observacao = 'fora do catalogo',
                            responsavel = 'Auditoria'
                    """)

                checks = {
                    check["verificacao"]: check["resultado"]
                    for check in db_module.quality_checks(catalog)
                }
                self.assertEqual(
                    checks["Produto desconhecido contendo alias oficial"], 0
                )
                self.assertEqual(
                    checks["Revisões automáticas Não incluir com preço alto"],
                    1,
                )

    def test_quality_checks_report_duplicates_without_excerpt(self):
        with tempfile.TemporaryDirectory() as directory:
            database_path = os.path.join(directory, "corrected.db")
            with patch.object(
                db_module, "DB_PATH", database_path
            ):
                db_module.initialize_database()
                db_module.upsert_report({
                    "id_relatorio": 1,
                    "url_solicitada": "https://example.test/1",
                    "url_final": "https://example.test/1",
                    "codigo_http": 200,
                    "data_relatorio": "2026-06-09",
                    "titulo": "Teste",
                    "texto_completo": "texto",
                    "html_bruto": "<p>teste</p>",
                    "hash_conteudo": "abc",
                    "eh_mercado": True,
                })
                base_price = {
                    "data_relatorio": "2026-06-09",
                    "produto_chines": "阿维菌素",
                    "produto_ingles": "Abamectin",
                    "produto_portugues": "ABAMECTINA",
                    "categoria": "inseticida",
                    "concentracao": "",
                    "forma_preco": "po_cristalino",
                    "forma_original": "阿维菌素精粉",
                    "segmento_preco": "nao_identificado",
                    "preco_yuan_ton": 690000,
                    "unidade_original": "69万元/吨",
                    "url": "https://example.test/1",
                    "trecho_original": "trecho A",
                    "nome_encontrado": "阿维菌素",
                    "posicao_texto": 10,
                    "regra_aplicada": "cotacao_generica",
                    "confianca": "alta",
                    "extractor_version": "teste",
                    "catalog_version": "teste",
                }
                second_price = {**base_price, "trecho_original": "trecho B", "posicao_texto": 20}
                db_module.replace_report_results(
                    1, [base_price, second_price], [], "extraido"
                )
                checks = {
                    check["verificacao"]: check
                    for check in db_module.quality_checks()
                }
                self.assertEqual(
                    checks["Duplicidade operacional sem considerar trecho"]["resultado"],
                    1,
                )
                self.assertEqual(
                    checks["Duplicidade operacional sem considerar trecho"]["status"],
                    "MONITORAR",
                )

    def test_deduplicates_only_automatic_rejected_reviews(self):
        with tempfile.TemporaryDirectory() as directory:
            database_path = os.path.join(directory, "corrected.db")
            with patch.object(
                db_module, "DB_PATH", database_path
            ):
                db_module.initialize_database()
                db_module.upsert_report({
                    "id_relatorio": 1,
                    "url_solicitada": "https://example.test/1",
                    "url_final": "https://example.test/1",
                    "codigo_http": 200,
                    "data_relatorio": "2026-06-09",
                    "titulo": "Teste",
                    "texto_completo": "texto",
                    "html_bruto": "<p>teste</p>",
                    "hash_conteudo": "abc",
                    "eh_mercado": True,
                })
                with db_module.connection() as connection:
                    for responsible in (
                        "Regra automática - catálogo oficial",
                        "Regra automática - catálogo oficial",
                        "Auditoria humana",
                    ):
                        connection.execute("""
                            INSERT INTO revisao_manual (
                                id_relatorio, data_relatorio, produto_sugerido,
                                preco_yuan_ton, trecho_original, motivo,
                                confianca, decisao, observacao, responsavel
                            ) VALUES (
                                1, '2026-06-09', 'produto_desconhecido',
                                150000, 'produto fora catalogo',
                                'teste', 'baixa', 'Não incluir',
                                'ok', ?
                            )
                        """, (responsible,))
                removed = db_module.deduplicate_auto_rejected_reviews()
                reviews = db_module.fetch_reviews()
                self.assertEqual(removed, 1)
                self.assertEqual(len(reviews), 2)
                self.assertEqual(
                    sum(1 for review in reviews if review[10] == "Auditoria humana"),
                    1,
                )


class ReviewImporterTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.catalog = load_catalog()

    def _review_workbook(self, decision, observation, responsible):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Pendentes"
        sheet.append([
            "ID do Caso", "ID Relatório", "Data", "Produto Sugerido",
            "Preço (yuan/ton)", "Trecho Original", "Motivo", "Confiança",
            "Decisão", "Observação", "Responsável",
        ])
        sheet.append([
            1, 10, "2026-06-09", "Glufosinato de amônio", 50000,
            "trecho", "motivo", "media", decision,
            observation, responsible,
        ])
        return workbook

    def test_reads_valid_include_decision(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "review.xlsx")
            workbook = self._review_workbook(
                "Incluir", "Conferido no relatório", "Analista"
            )
            workbook.save(path)
            workbook.close()
            decisions = read_manual_decisions(self.catalog, path)
            self.assertEqual(decisions[0]["decision"], "Incluir")
            self.assertEqual(
                decisions[0]["product"]["ingles"],
                "Glufosinate-ammonium",
            )

    def test_rejects_decision_without_responsible(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "review.xlsx")
            workbook = self._review_workbook(
                "Não incluir", "Intermediário", ""
            )
            workbook.save(path)
            workbook.close()
            with self.assertRaises(ReviewImportError):
                read_manual_decisions(self.catalog, path)

    def test_correction_uses_product_from_observation(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "review.xlsx")
            workbook = self._review_workbook(
                "Corrigir", "Produto correto: L-GLUFOSINATO", "Analista"
            )
            workbook.save(path)
            workbook.close()
            decisions = read_manual_decisions(self.catalog, path)
            self.assertEqual(
                decisions[0]["product"]["ingles"], "L-Glufosinate"
            )


class ExporterTests(unittest.TestCase):
    def test_review_has_decision_dropdown_for_future_rows(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            path = os.path.join(directory, "review.xlsx")
            with patch.object(
                exporter_module, "fetch_reviews", return_value=[]
            ), patch.object(
                exporter_module, "quality_checks", return_value=[]
            ), patch.object(
                exporter_module, "fetch_metadata", return_value={}
            ):
                exporter_module.export_review(
                    load_catalog(),
                    allow_review_overwrite=True,
                    target_path=path,
                )

            workbook = load_workbook(path, read_only=False, data_only=False)
            validations = list(
                workbook["Pendentes"].data_validations.dataValidation
            )
            workbook.close()
            self.assertEqual(len(validations), 1)
            self.assertEqual(
                validations[0].formula1,
                '"Incluir,Não incluir,Corrigir"',
            )
            self.assertEqual(str(validations[0].sqref), "I2:I10000")
            self.assertTrue(validations[0].showErrorMessage)

    def test_batch_publication_restores_files_after_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            staged = []
            targets = []
            for index in range(3):
                staged_path = os.path.join(
                    directory, f"novo_{index}.xlsx"
                )
                target_path = os.path.join(
                    directory, f"saida_{index}.xlsx"
                )
                with open(staged_path, "w", encoding="utf-8") as file:
                    file.write(f"novo {index}")
                with open(target_path, "w", encoding="utf-8") as file:
                    file.write(f"antigo {index}")
                staged.append(staged_path)
                targets.append(target_path)

            real_replace = os.replace

            def fail_on_second_file(source, target):
                if target == targets[1]:
                    raise PermissionError("arquivo aberto")
                real_replace(source, target)

            with patch.object(
                exporter_module.os,
                "replace",
                side_effect=fail_on_second_file,
            ):
                with self.assertRaises(PermissionError):
                    exporter_module._publish_staged_files(
                        list(zip(staged, targets))
                    )

            for index, target in enumerate(targets):
                with open(target, encoding="utf-8") as file:
                    self.assertEqual(file.read(), f"antigo {index}")


if __name__ == "__main__":
    unittest.main()
