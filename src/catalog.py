"""Leitura e validação do catálogo oficial de produtos.

A planilha de catálogo é a fonte de verdade. O sistema para
com uma mensagem clara se a planilha estiver incompleta ou contraditória,
evitando publicar preços com traduções duvidosas.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
import hashlib
import json
import os
from typing import Iterable

from openpyxl import load_workbook

from config.settings import ALIASES_PATH, CATALOG_PATH


class CatalogValidationError(ValueError):
    """Erro de catálogo que pode ser apresentado diretamente ao usuário."""


@dataclass(frozen=True)
class Catalog:
    products: dict[str, dict]
    aliases: dict[str, str]
    file_hash: str
    file_modified_at: str

    @property
    def version(self) -> str:
        """Versão curta e estável baseada no conteúdo da planilha."""
        return self.file_hash[:12]


def _normalize_cell(value) -> str:
    return str(value).strip() if value is not None else ""


def _category_for_sheet(sheet_name: str, row: tuple) -> str:
    fixed = {
        "Herbicidas": "herbicida",
        "Inseticidas": "inseticida",
        "Fungicidas": "fungicida",
    }
    if sheet_name in fixed:
        return fixed[sheet_name]
    category = _normalize_cell(row[3] if len(row) > 3 else "").lower()
    aliases = {
        "herbicida": "herbicida",
        "inseticida": "inseticida",
        "fungicida": "fungicida",
    }
    return aliases.get(category, "")


def _iter_catalog_rows(workbook) -> Iterable[tuple[str, int, tuple]]:
    """Entrega somente as linhas de dados, ignorando títulos explicativos."""
    for worksheet in workbook.worksheets:
        header_row = None
        for row_number, row in enumerate(
            worksheet.iter_rows(values_only=True), start=1
        ):
            values = tuple(_normalize_cell(value) for value in row)
            if "Português" in values and "Inglês" in values and "Chinês" in values:
                header_row = row_number
                continue
            if header_row and row_number > header_row and any(values[:3]):
                yield worksheet.title, row_number, values


def load_catalog(path: str = CATALOG_PATH) -> Catalog:
    """Carrega a planilha oficial ou interrompe com orientação objetiva."""
    if not os.path.exists(path):
        raise CatalogValidationError(
            f"Catálogo não encontrado em '{path}'. Ajuste CATALOG_PATH "
            "ou use catalog/example_catalog.xlsx."
        )

    with open(path, "rb") as file:
        file_hash = hashlib.sha256(file.read()).hexdigest()

    workbook = load_workbook(path, read_only=True, data_only=True)
    required_sheets = {
        "Herbicidas", "Inseticidas", "Fungicidas", "Outros do Banco",
    }
    missing_sheets = required_sheets - set(workbook.sheetnames)
    if missing_sheets:
        workbook.close()
        raise CatalogValidationError(
            "A planilha oficial está sem a(s) aba(s): "
            + ", ".join(sorted(missing_sheets))
            + "."
        )
    products: dict[str, dict] = {}
    try:
        for sheet, row_number, row in _iter_catalog_rows(workbook):
            portuguese = _normalize_cell(row[0] if len(row) > 0 else "")
            english = _normalize_cell(row[1] if len(row) > 1 else "")
            chinese = _normalize_cell(row[2] if len(row) > 2 else "")
            category = _category_for_sheet(sheet, row)

            missing = [
                label for label, value in (
                    ("Português", portuguese),
                    ("Inglês", english),
                    ("Chinês", chinese),
                    ("Categoria", category),
                ) if not value
            ]
            if missing:
                raise CatalogValidationError(
                    f"Catálogo incompleto na aba '{sheet}', linha {row_number}: "
                    f"preencha {', '.join(missing)}."
                )

            new_value = {
                "ingles": english,
                "portugues": portuguese,
                "categoria": category,
                "aliases": [],
            }
            previous = products.get(chinese)
            if previous and previous != new_value:
                raise CatalogValidationError(
                    f"O nome chinês '{chinese}' aparece com relações diferentes "
                    f"na aba '{sheet}', linha {row_number}."
                )
            if previous:
                raise CatalogValidationError(
                    f"O nome chinês '{chinese}' está duplicado na aba '{sheet}', "
                    f"linha {row_number}."
                )
            products[chinese] = new_value
    finally:
        workbook.close()

    if not products:
        raise CatalogValidationError(
            "A planilha oficial não contém produtos válidos."
        )

    # Aliases opcionais vêm de um JSON {"alias": "nome principal"}. Só valem
    # se o produto principal existir no catálogo; a tradução vem sempre da
    # planilha.
    aliases = {name: name for name in products}
    if os.path.exists(ALIASES_PATH):
        with open(ALIASES_PATH, encoding="utf-8") as file:
            extra = json.load(file)
        for alias, main_name in extra.items():
            if main_name not in products:
                continue
            owner = aliases.get(alias)
            if owner and owner != main_name:
                raise CatalogValidationError(
                    f"O alias chinês '{alias}' aponta para mais de um produto."
                )
            aliases[alias] = main_name

    modified = datetime.fromtimestamp(os.path.getmtime(path)).isoformat(
        timespec="seconds"
    )
    return Catalog(
        products=products,
        aliases=aliases,
        file_hash=file_hash,
        file_modified_at=modified,
    )
