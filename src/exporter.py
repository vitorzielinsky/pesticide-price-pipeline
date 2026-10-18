"""Geração dos Excels (preços e revisão manual) com publicação atômica."""

from __future__ import annotations

from datetime import datetime
import os
import shutil
import tempfile
import uuid

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill

from config.settings import (
    EXCEL_PATH,
    EXPORTS_DIR,
    REVIEW_EXCEL_PATH,
)
from src.catalog import Catalog
from src.database import (
    fetch_metadata,
    fetch_prices,
    fetch_reports,
    fetch_reviews,
    quality_checks,
)
COR_CABECALHO = "FF1F6B42"
COR_FONTE_CABECALHO = "FFFFFFFF"
COR_HERBICIDA = "FFE8F5E9"
COR_INSETICIDA = "FFFBE9E7"
COR_FUNGICIDA = "FFE3F2FD"


HEADER_FILL = PatternFill("solid", fgColor=COR_CABECALHO)
HEADER_FONT = Font(bold=True, color=COR_FONTE_CABECALHO)
CATEGORY_COLORS = {
    "herbicida": COR_HERBICIDA,
    "inseticida": COR_INSETICIDA,
    "fungicida": COR_FUNGICIDA,
}


def _style_header(worksheet, headers: list[str]) -> None:
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    worksheet.freeze_panes = "A2"


def _set_widths(worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[
            openpyxl.utils.get_column_letter(index)
        ].width = width


def _atomic_save(workbook, target_path: str) -> None:
    """Publica somente um arquivo completo e preserva a versão anterior."""
    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    fd, temporary_path = tempfile.mkstemp(
        suffix=".xlsx", dir=os.path.dirname(target_path)
    )
    os.close(fd)
    try:
        workbook.save(temporary_path)
        if os.path.exists(target_path):
            shutil.copy2(target_path, f"{target_path}.bak")
        os.replace(temporary_path, target_path)
    finally:
        if os.path.exists(temporary_path):
            os.remove(temporary_path)


def export_operational(
    catalog: Catalog,
    target_path: str = EXCEL_PATH,
) -> str:
    """Mantem as abas atuais e adiciona forma/base ao fim da aba Precos."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    prices_sheet = workbook.create_sheet("Preços")
    price_headers = [
        "Data do Relatório", "Produto (Inglês)", "Produto (Português)",
        "Categoria", "Concentração", "Preço (yuan/ton)",
        "Produto (Chinês)", "ID Relatório", "URL",
        "Forma/Base", "Forma original", "Segmento",
    ]
    _style_header(prices_sheet, price_headers)
    prices = fetch_prices(high_confidence_only=True)
    for row_number, row in enumerate(prices, start=2):
        (
            data, report_id, chinese, english, portuguese, category,
            concentration, form_price, form_original, segment, price, _, url,
            _, _, _,
        ) = row
        values = [
            data, english, portuguese, category, concentration or None,
            price, chinese, report_id, url, form_price,
            form_original or None, segment,
        ]
        prices_sheet.append(values)
        fill = PatternFill(
            "solid", fgColor=CATEGORY_COLORS.get(category, "FFFFFFFF")
        )
        for cell in prices_sheet[row_number]:
            cell.fill = fill
            cell.alignment = Alignment(vertical="center")
        prices_sheet.cell(row_number, 6).number_format = "#,##0.00"
    _set_widths(
        prices_sheet, [16, 30, 35, 14, 16, 18, 20, 14, 60, 20, 32, 24]
    )
    prices_sheet.auto_filter.ref = f"A1:L{max(1, prices_sheet.max_row)}"

    reports_sheet = workbook.create_sheet("Relatórios")
    report_headers = [
        "ID", "Data do Relatório", "Título", "É Mercado?",
        "Processado?", "URL", "Data de Coleta",
    ]
    _style_header(reports_sheet, report_headers)
    for report in fetch_reports():
        report_id, date, title, market, state, url, collected_at, _ = report
        reports_sheet.append([
            report_id, date, title, "Sim" if market else "Não",
            "Sim" if state in {"extraido", "sem_preco_confirmado"} else "Não",
            url, collected_at,
        ])
    _set_widths(reports_sheet, [10, 16, 60, 12, 14, 60, 20])
    reports_sheet.auto_filter.ref = f"A1:G{max(1, reports_sheet.max_row)}"

    products_sheet = workbook.create_sheet("Produtos")
    _style_header(products_sheet, [
        "Produto (Chinês)", "Produto (Inglês)",
        "Produto (Português)", "Categoria",
    ])
    for chinese, product in sorted(
        catalog.products.items(),
        key=lambda item: (item[1]["categoria"], item[1]["ingles"]),
    ):
        products_sheet.append([
            chinese, product["ingles"], product["portugues"],
            product["categoria"].capitalize(),
        ])
    _set_widths(products_sheet, [20, 35, 40, 14])

    workbook.properties.title = "Pesticide prices"
    workbook.properties.description = (
        f"Gerado em {datetime.now():%Y-%m-%d %H:%M}; "
        f"catálogo {catalog.version}"
    )
    _atomic_save(workbook, target_path)
    return target_path


def _has_unimported_review_decisions(path: str) -> bool:
    """Evita apagar decisões digitadas antes que o usuário as incorpore."""
    if not os.path.exists(path):
        return False
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "Pendentes" not in workbook.sheetnames:
            return False
        sheet = workbook["Pendentes"]
        headers = {
            str(cell.value).strip(): cell.column
            for cell in sheet[1]
            if cell.value is not None
        }
        decision_column = headers.get("Decisão")
        if decision_column is None:
            return False
        return any(
            str(sheet.cell(row, decision_column).value or "").strip()
            for row in range(2, sheet.max_row + 1)
        )
    finally:
        workbook.close()


def export_review(
    catalog: Catalog,
    allow_review_overwrite: bool = False,
    target_path: str = REVIEW_EXCEL_PATH,
) -> str:
    if (
        not allow_review_overwrite
        and _has_unimported_review_decisions(REVIEW_EXCEL_PATH)
    ):
        raise ValueError(
            "Há decisões preenchidas no Excel de revisão. Execute "
            "import_review.py antes de atualizar ou coletar."
        )

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Pendentes"
    _style_header(sheet, [
        "ID do Caso", "ID Relatório", "Data", "Produto Sugerido",
        "Preço (yuan/ton)", "Trecho Original", "Motivo", "Confiança",
        "Decisão", "Observação", "Responsável",
    ])
    reviews = fetch_reviews()
    for row_number, row in enumerate(
        (item for item in reviews if not item[8]),
        start=2,
    ):
        values = list(row)
        if values[3] == "produto_desconhecido":
            values[3] = "Produto fora do catálogo"
        sheet.append(values)
        sheet.row_dimensions[row_number].height = 60
        for cell in sheet[row_number]:
            cell.alignment = Alignment(vertical="top")
        for column in (6, 7, 10, 11):
            sheet.cell(row_number, column).alignment = Alignment(
                vertical="top", wrap_text=True
            )
        sheet.cell(row_number, 5).number_format = "#,##0.00"
    _set_widths(sheet, [12, 14, 14, 28, 18, 90, 55, 12, 18, 35, 22])
    sheet.auto_filter.ref = f"A1:K{max(1, sheet.max_row)}"
    sheet.freeze_panes = "A2"
    decision_validation = DataValidation(
        type="list",
        formula1='"Incluir,Não incluir,Corrigir"',
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="stop",
        showInputMessage=True,
    )
    decision_validation.error = (
        "Escolha Incluir, Não incluir ou Corrigir."
    )
    decision_validation.errorTitle = "Decisão inválida"
    decision_validation.prompt = (
        "Selecione uma das decisões aceitas pelo sistema."
    )
    decision_validation.promptTitle = "Decisão"
    sheet.add_data_validation(decision_validation)
    # A faixa ampla mantém a lista disponível mesmo quando a aba é gerada
    # vazia e recebe novos casos em uma execução posterior.
    decision_validation.add("I2:I10000")

    resolved_sheet = workbook.create_sheet("Resolvidos")
    _style_header(resolved_sheet, [
        "ID do Caso", "ID Relatório", "Data", "Produto Sugerido",
        "Preço (yuan/ton)", "Trecho Original", "Motivo", "Confiança",
        "Decisão", "Observação", "Responsável",
    ])
    for row_number, row in enumerate(
        (item for item in reviews if item[8]),
        start=2,
    ):
        values = list(row)
        if values[3] == "produto_desconhecido":
            values[3] = "Produto fora do catálogo"
        resolved_sheet.append(values)
        resolved_sheet.row_dimensions[row_number].height = 42
        for cell in resolved_sheet[row_number]:
            cell.alignment = Alignment(vertical="top")
        for column in (6, 7, 10, 11):
            resolved_sheet.cell(row_number, column).alignment = Alignment(
                vertical="top", wrap_text=True
            )
        resolved_sheet.cell(row_number, 5).number_format = "#,##0.00"
    _set_widths(
        resolved_sheet, [12, 14, 14, 28, 18, 90, 55, 12, 18, 50, 32]
    )
    resolved_sheet.auto_filter.ref = (
        f"A1:K{max(1, resolved_sheet.max_row)}"
    )
    resolved_sheet.freeze_panes = "A2"

    checks_sheet = workbook.create_sheet("Verificações")
    _style_header(checks_sheet, ["Verificação", "Resultado", "Status"])
    for check in quality_checks(catalog):
        checks_sheet.append([
            check["verificacao"], check["resultado"], check["status"]
        ])
    _set_widths(checks_sheet, [55, 80, 14])


    _atomic_save(workbook, target_path)
    return target_path


def _publish_staged_files(staged_targets: list[tuple[str, str]]) -> None:
    """Substitui todas as saídas como um conjunto ou restaura as anteriores."""
    if not staged_targets:
        return
    rollback_directory = tempfile.mkdtemp(
        prefix=".rollback_",
        dir=os.path.dirname(staged_targets[0][1]),
    )
    originals: dict[str, str | None] = {}
    replaced: list[str] = []
    try:
        # Primeiro preservamos todas as saídas. Nenhum arquivo real é
        # substituído antes de os novos Excels e os backups existirem.
        for index, (_, target) in enumerate(staged_targets):
            if os.path.exists(target):
                rollback_path = os.path.join(
                    rollback_directory, f"original_{index}.xlsx"
                )
                shutil.copy2(target, rollback_path)
                shutil.copy2(target, f"{target}.bak")
                originals[target] = rollback_path
            else:
                originals[target] = None

        try:
            for staged, target in staged_targets:
                os.replace(staged, target)
                replaced.append(target)
        except Exception:
            # Se uma substituição falhar, os arquivos já trocados voltam ao
            # estado anterior. Assim ninguém recebe um conjunto misto de arquivos.
            for target in replaced:
                original = originals[target]
                if original is None:
                    if os.path.exists(target):
                        os.remove(target)
                else:
                    shutil.copy2(original, target)
            raise
    finally:
        shutil.rmtree(rollback_directory, ignore_errors=True)


def export_all(
    catalog: Catalog,
    allow_review_overwrite: bool = False,
) -> list[str]:
    # A verificação vem antes de qualquer saída para evitar atualização parcial.
    if (
        not allow_review_overwrite
        and _has_unimported_review_decisions(REVIEW_EXCEL_PATH)
    ):
        raise ValueError(
            "Há decisões preenchidas no Excel de revisão. Execute "
            "import_review.py antes de atualizar ou coletar."
        )
    os.makedirs(EXPORTS_DIR, exist_ok=True)
    # Arquivos preparados na própria pasta exports evitam movimentações entre
    # diretórios (pastas sincronizadas na nuvem podem bloquear arquivos).
    staging_token = uuid.uuid4().hex
    targets = [
        EXCEL_PATH,
        REVIEW_EXCEL_PATH,
    ]
    staged = [
        os.path.join(
            EXPORTS_DIR,
            f".staging_{staging_token}_{os.path.basename(path)}",
        )
        for path in targets
    ]
    try:
        export_operational(catalog, target_path=staged[0])
        export_review(
            catalog,
            allow_review_overwrite=allow_review_overwrite,
            target_path=staged[1],
        )
        _publish_staged_files(list(zip(staged, targets)))
    finally:
        for staged_path in staged:
            if os.path.exists(staged_path):
                os.remove(staged_path)
    return targets
