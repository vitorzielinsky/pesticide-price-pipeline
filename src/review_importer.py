"""Importa decisões preenchidas no Excel de revisão."""

from __future__ import annotations

import os
import re
import shutil
from datetime import datetime
import unicodedata

from openpyxl import load_workbook

from config.settings import DB_PATH, REVIEW_EXCEL_PATH
from src.catalog import Catalog
from src.database import apply_manual_review_decisions


class ReviewImportError(ValueError):
    """Mensagem corrigível que pode ser apresentada ao usuário."""


def _key(value) -> str:
    text = str(value or "").strip().casefold()
    text = unicodedata.normalize("NFKD", text)
    return "".join(char for char in text if not unicodedata.combining(char))


def _cell_text(value) -> str:
    """Converte células vazias em texto vazio, sem gerar a palavra 'None'."""
    return "" if value is None else str(value).strip()


def _product_index(catalog: Catalog) -> dict[str, dict]:
    index = {}
    for chinese, product in catalog.products.items():
        value = {
            "chines": chinese,
            "ingles": product["ingles"],
            "portugues": product["portugues"],
            "categoria": product["categoria"],
        }
        for name in (chinese, product["ingles"], product["portugues"]):
            normalized = _key(name)
            previous = index.get(normalized)
            if previous and previous["chines"] != chinese:
                continue
            index[normalized] = value
    return index


def _corrected_product_text(observation: str) -> str:
    text = observation.strip()
    match = re.search(
        r"(?:produto\s+correto|corrigir\s+para)\s*[:=-]\s*(.+)",
        text,
        flags=re.IGNORECASE,
    )
    return match.group(1).strip() if match else text


def read_manual_decisions(
    catalog: Catalog,
    path: str = REVIEW_EXCEL_PATH,
) -> list[dict]:
    if not os.path.exists(path):
        raise ReviewImportError(
            "O arquivo manual_review.xlsx não foi encontrado."
        )
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except PermissionError as error:
        raise ReviewImportError(
            "Feche o Excel de revisão antes de incorporar as decisões."
        ) from error

    try:
        if "Pendentes" not in workbook.sheetnames:
            raise ReviewImportError(
                "A aba 'Pendentes' não foi encontrada no Excel de revisão."
            )
        sheet = workbook["Pendentes"]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        return []
    headers = {
        str(value).strip(): index
        for index, value in enumerate(rows[0])
        if value is not None
    }
    required = {
        "ID do Caso", "Produto Sugerido", "Decisão",
        "Observação", "Responsável",
    }
    missing = required - set(headers)
    if missing:
        raise ReviewImportError(
            "Colunas obrigatórias ausentes: " + ", ".join(sorted(missing))
        )

    valid_decisions = {
        "incluir": "Incluir",
        "nao incluir": "Não incluir",
        "não incluir": "Não incluir",
        "corrigir": "Corrigir",
    }
    products = _product_index(catalog)
    decisions = []

    for row_number, row in enumerate(rows[1:], start=2):
        raw_decision = row[headers["Decisão"]] if len(row) > headers["Decisão"] else None
        if not str(raw_decision or "").strip():
            continue
        decision = valid_decisions.get(_key(raw_decision))
        if not decision:
            raise ReviewImportError(
                f"Linha {row_number}: use Incluir, Não incluir ou Corrigir."
            )
        responsible = _cell_text(
            row[headers["Responsável"]]
            if len(row) > headers["Responsável"] else ""
        )
        if not responsible:
            raise ReviewImportError(
                f"Linha {row_number}: preencha o Responsável."
            )
        observation = _cell_text(
            row[headers["Observação"]]
            if len(row) > headers["Observação"] else ""
        )
        suggested = _cell_text(
            row[headers["Produto Sugerido"]]
            if len(row) > headers["Produto Sugerido"] else ""
        )

        product = None
        product_text = suggested
        if decision == "Incluir":
            product = products.get(_key(suggested))
            if product is None:
                raise ReviewImportError(
                    f"Linha {row_number}: o produto sugerido não foi encontrado "
                    "no catálogo oficial. Use Corrigir."
                )
        elif decision == "Corrigir":
            if not observation:
                raise ReviewImportError(
                    f"Linha {row_number}: informe o produto correto na Observação."
                )
            product_text = _corrected_product_text(observation)
            product = products.get(_key(product_text))
            if product is None:
                raise ReviewImportError(
                    f"Linha {row_number}: produto correto '{product_text}' não "
                    "foi encontrado no catálogo oficial."
                )

        review_id = row[headers["ID do Caso"]]
        try:
            review_id = int(review_id)
        except (TypeError, ValueError) as error:
            raise ReviewImportError(
                f"Linha {row_number}: ID do caso inválido."
            ) from error

        decisions.append({
            "review_id": review_id,
            "decision": decision,
            "observation": observation,
            "responsible": responsible,
            "product": product,
            "product_text": product_text,
        })
    return decisions


def import_manual_reviews(catalog: Catalog) -> dict:
    decisions = read_manual_decisions(catalog)
    if not decisions:
        return {
            "total": 0, "incluidos": 0,
            "nao_incluidos": 0, "corrigidos": 0,
        }

    backup = (
        f"{DB_PATH}.antes_revisao_"
        f"{datetime.now():%Y%m%d_%H%M%S}.bak"
    )
    shutil.copy2(DB_PATH, backup)
    result = apply_manual_review_decisions(decisions, catalog.version)
    result["backup"] = backup
    return result
